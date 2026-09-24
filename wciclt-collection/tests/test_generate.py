from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from wciclt.generate import discover_categories, generate_week, sync_categories_across
from wciclt.models import Service, ServiceType, WeekPlan


@pytest.fixture
def tmp_paths(tmp_path: Path):
    analysis = tmp_path / "2026-Collection Analysis.xlsx"
    output = tmp_path / "reports"
    output.mkdir()
    return analysis, output


def _week_two_services() -> WeekPlan:
    return WeekPlan(
        week_end=date(2026, 9, 27),
        services=[
            Service(date(2026, 9, 23), ServiceType.MIDWEEK, recorded_by="Bro Ernest, Sis Sandra"),
            Service(date(2026, 9, 27), ServiceType.SUNDAY, recorded_by="Bro Ernest, Sis Sandra"),
        ],
    )


def test_other_requires_details():
    plan = WeekPlan(
        week_end=date(2026, 9, 27),
        services=[Service(date(2026, 9, 24), ServiceType.OTHER)],
    )
    with pytest.raises(ValueError, match="details"):
        plan.validate()


def test_generate_two_service_week(tmp_paths):
    analysis, output = tmp_paths
    result = generate_week(_week_two_services(), analysis, output)
    assert result.tabs == ["Wed-09.23.2026", "Sun-09.27.2026"]
    assert result.report_path and result.report_path.exists()

    wb = load_workbook(analysis)
    assert "_Categories" in wb.sheetnames
    assert wb["_Categories"].sheet_state == "hidden"
    sun = wb["Sun-09.27.2026"]
    assert sun["C2"].value == "09/27/2026 - Sunday Service"
    assert sun["B1"].value == "WINNERS CHAPEL CARLOTTE"
    assert sun["B2"].value == "COLLECTIONS/DEPOSIT RECON"
    assert sun["K9"].value == "=SUM(I9*J9)"
    assert sun["E32"].value == "=SUM(E7:E31)"
    assert sun["B35"].value == "Offering"
    assert "TOTAL COLLECTIONS ANALYSIS" in [sun.cell(r, 2).value for r in range(35, 60)]
    wb.close()

    report = load_workbook(result.report_path)
    ws = report.active
    assert ws.title == "09.27.26"
    assert "Collection Report as at" in ws["A2"].value
    assert ws["A4"].value == "Date:09.23.2026"
    assert ws["B4"].value == "Midweek Service"
    assert ws["D4"].value == "Date:09.27.2026"
    assert ws["E4"].value == "Sunday Service"
    assert ws["B12"].value == "=SUM(B7:B9)"
    assert ws["A15"].value == "Offering"
    assert ws["D15"].value == "=A15"
    last_cat = 14 + len(result.categories)
    assert ws.cell(last_cat + 1, 1).value == "TOTAL"
    assert "SUM(B15:" in str(ws.cell(last_cat + 1, 2).value)
    report.close()


def test_generate_variable_service_count(tmp_paths):
    analysis, output = tmp_paths
    plan = WeekPlan(
        week_end=date(2026, 9, 6),
        services=[
            Service(date(2026, 9, 2), ServiceType.WOSE_DAY_1),
            Service(date(2026, 9, 3), ServiceType.WOSE_DAY_2),
            Service(date(2026, 9, 4), ServiceType.WOSE_DAY_3),
            Service(date(2026, 9, 6), ServiceType.SUNDAY),
            Service(date(2026, 9, 5), ServiceType.OTHER, details="Youth Night"),
        ],
    )
    result = generate_week(plan, analysis, output)
    assert len(result.tabs) == 5
    assert "Fri-09.04.2026" in result.tabs
    wb = load_workbook(result.report_path)
    ws = wb.active
    labels = [ws.cell(4, 2 + i * 3).value for i in range(5)]
    assert labels == [
        "WOSE Day 1",
        "WOSE Day 2",
        "WOSE Day 3",
        "Sunday Service",
        "Youth Night",
    ]
    wb.close()


def test_category_sync_add_remove(tmp_paths):
    analysis, output = tmp_paths
    generate_week(_week_two_services(), analysis, output)
    cats = discover_categories(analysis)
    cats.append("Mission Seed")
    cats = [c for c in cats if c != "Chairs"]
    sync_categories_across(analysis, cats)
    wb = load_workbook(analysis)
    names = [wb["_Categories"].cell(r, 1).value for r in range(2, 30) if wb["_Categories"].cell(r, 1).value]
    assert "Mission Seed" in names
    assert "Chairs" not in names
    sun = wb["Sun-09.27.2026"]
    col_b = [sun.cell(r, 2).value for r in range(35, 55)]
    assert "Mission Seed" in col_b
    assert "Chairs" not in col_b
    wb.close()


def test_cdot_and_shiloh_labels(tmp_paths):
    analysis, output = tmp_paths
    plan = WeekPlan(
        week_end=date(2026, 12, 14),
        services=[
            Service(date(2026, 12, 9), ServiceType.SHILOH_OPENING),
            Service(date(2026, 12, 10), ServiceType.SHILOH_HOUR),
            Service(date(2026, 12, 11), ServiceType.SHILOH_ENCOUNTER),
            Service(date(2026, 12, 12), ServiceType.SHILOH_IMPARTATION),
            Service(date(2026, 12, 14), ServiceType.CDOT),
        ],
    )
    result = generate_week(plan, analysis, output)
    wb = load_workbook(analysis)
    assert wb["Wed-12.09.2026"]["C2"].value == "12/09/2026 - Shiloh Opening Night"
    assert wb["Mon-12.14.2026"]["C2"].value == "12/14/2026 - CDOT"
    wb.close()
    assert result.report_path.exists()
