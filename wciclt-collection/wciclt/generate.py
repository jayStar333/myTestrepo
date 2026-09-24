from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .analysis import add_service_tab, apply_categories_to_sheet, create_or_open_analysis
from .categories import (
    infer_from_analysis_workbook,
    infer_from_report_workbook,
    load_json_categories,
    merge_unique,
    resolve_categories,
    save_json_categories,
    write_categories_sheet,
)
from .models import WeekPlan
from .paths import DEFAULT_CATEGORIES_JSON
from .report import build_weekly_report


@dataclass
class GenerateResult:
    analysis_path: Path
    report_path: Path | None
    tabs: list[str]
    categories: list[str]


def sync_categories_across(
    analysis_path: Path,
    categories: list[str],
    report_path: Path | None = None,
    json_path: Path | None = None,
) -> list[str]:
    save_json_categories(categories, json_path or DEFAULT_CATEGORIES_JSON)
    analysis_path.parent.mkdir(parents=True, exist_ok=True)
    wb = create_or_open_analysis(analysis_path)
    try:
        write_categories_sheet(wb, categories)
        for name in wb.sheetnames:
            if name.startswith("_"):
                continue
            apply_categories_to_sheet(wb[name], categories)
        wb.save(analysis_path)
    finally:
        wb.close()
    return categories


def discover_categories(
    analysis_path: Path | None,
    report_path: Path | None = None,
    explicit: list[str] | None = None,
) -> list[str]:
    if explicit:
        return resolve_categories(explicit=explicit)
    found: list[list[str]] = []
    if analysis_path and analysis_path.exists():
        wb = load_workbook(analysis_path)
        try:
            from_analysis = infer_from_analysis_workbook(wb)
            if from_analysis:
                found.append(from_analysis)
        finally:
            wb.close()
    if report_path and report_path.exists():
        from_report = infer_from_report_workbook(report_path)
        if from_report:
            found.append(from_report)
    if found:
        return merge_unique(*found)
    return load_json_categories()


def generate_week(
    plan: WeekPlan,
    analysis_path: Path,
    output_dir: Path,
    categories: list[str] | None = None,
    write_report: bool = True,
    write_analysis: bool = True,
    sync_existing_tabs: bool = False,
    json_path: Path | None = None,
) -> GenerateResult:
    plan.validate()
    cats = categories or discover_categories(analysis_path)
    save_json_categories(cats, json_path or DEFAULT_CATEGORIES_JSON)
    tabs: list[str] = []
    report_path = None

    if write_analysis:
        analysis_path.parent.mkdir(parents=True, exist_ok=True)
        wb = create_or_open_analysis(analysis_path)
        try:
            write_categories_sheet(wb, cats)
            if sync_existing_tabs:
                for name in wb.sheetnames:
                    if name.startswith("_"):
                        continue
                    apply_categories_to_sheet(wb[name], cats)
            for service in plan.services:
                tabs.append(add_service_tab(wb, service, cats))
            wb.save(analysis_path)
        finally:
            wb.close()

    if write_report:
        report_path = output_dir / plan.report_filename()
        build_weekly_report(plan, cats, report_path)

    return GenerateResult(
        analysis_path=analysis_path,
        report_path=report_path,
        tabs=tabs,
        categories=cats,
    )
