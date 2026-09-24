from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import ATTENDANCE_ROWS, CURRENCY_FORMAT, Service, WeekPlan

THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
# ARGB (FF prefix). A 6-digit RGB becomes 00xxxxxx in openpyxl and is invisible.
YELLOW = PatternFill("solid", fgColor="FFFFFF00")
COLLECTION_FILL = PatternFill("solid", fgColor="FFC6EFCE")
# Source spacer: theme-3 (or 7) light tint + theme-1 (black) sideways text.
SPACER_FILLS = (
    PatternFill(fill_type="solid", fgColor=Color(theme=3, tint=0.5999938962981048)),
    PatternFill(fill_type="solid", fgColor=Color(theme=7, tint=0.5999938962981048)),
)
THEME_DARK = Color(theme=1)
TITLE_FONT = Font(name="Calibri", size=24, bold=True, color=THEME_DARK)
HEAD_FONT = Font(name="Calibri", size=18, bold=True, color=THEME_DARK)
SECTION_FONT = Font(name="Calibri", size=11, bold=True, color=THEME_DARK)
BODY_FONT = Font(name="Calibri", size=11, color=THEME_DARK)
TOTAL_ATT_FONT = Font(name="Calibri", size=14, bold=True, color=THEME_DARK)
TOTAL_COL_FONT = Font(name="Calibri", size=18, bold=True, color=THEME_DARK)
AMOUNT_FONT = Font(name="Cambria", size=14, color=THEME_DARK)
VERT_FONT = Font(name="Calibri", size=26, bold=True, color=THEME_DARK)


def _style_range(ws: Worksheet, row: int, col: int, font=None, fill=None, border=True, align=None, fmt=None):
    cell = ws.cell(row, col)
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if border:
        cell.border = BOX
    if align is not None:
        cell.alignment = align
    if fmt is not None:
        cell.number_format = fmt


def write_service_block(
    ws: Worksheet,
    service: Service,
    index: int,
    categories: list[str],
    first_label_col: int,
    amounts: list[float] | None = None,
) -> int:
    label_col = 1 + index * 3
    value_col = label_col + 1
    spacer_col = label_col + 2
    last_row = 14 + len(categories) + 1

    ws.column_dimensions[get_column_letter(label_col)].width = 26.3
    ws.column_dimensions[get_column_letter(value_col)].width = 22.0
    ws.column_dimensions[get_column_letter(spacer_col)].width = 6.3

    service_name = service.display_label()
    ws.cell(4, label_col, service.report_date_label())
    ws.cell(4, value_col, service_name)
    _style_range(ws, 4, label_col, HEAD_FONT, align=Alignment(vertical="center", wrap_text=True))
    _style_range(ws, 4, value_col, HEAD_FONT, align=Alignment(vertical="center", wrap_text=True))
    ws.row_dimensions[4].height = 117.0

    ws.cell(5, label_col, "Recorded By:")
    ws.cell(5, value_col, service.recorded_by or "")
    for col in (label_col, value_col):
        _style_range(ws, 5, col, SECTION_FONT, fill=YELLOW, align=Alignment(vertical="center", wrap_text=True))
    ws.row_dimensions[5].height = 36

    ws.cell(6, label_col, "ATTENDANCE")
    _style_range(ws, 6, label_col, SECTION_FONT, align=Alignment(horizontal="center", vertical="center"))
    _style_range(ws, 6, value_col, SECTION_FONT)

    for row, name in ATTENDANCE_ROWS:
        ws.cell(row, label_col, name)
        ws.cell(row, value_col, 0)
        _style_range(ws, row, label_col, BODY_FONT, align=Alignment(vertical="center"))
        _style_range(ws, row, value_col, BODY_FONT)

    ws.cell(12, label_col, "TOTAL")
    ws.cell(12, value_col, f"=SUM({get_column_letter(value_col)}7:{get_column_letter(value_col)}9)")
    _style_range(ws, 12, label_col, TOTAL_ATT_FONT, align=Alignment(vertical="center"))
    _style_range(ws, 12, value_col, TOTAL_ATT_FONT)

    ws.merge_cells(start_row=13, start_column=label_col, end_row=13, end_column=value_col)
    ws.cell(13, label_col, "COLLECTION")
    _style_range(ws, 13, label_col, SECTION_FONT, fill=COLLECTION_FILL, align=Alignment(horizontal="center", vertical="center"))
    _style_range(ws, 13, value_col, SECTION_FONT, fill=COLLECTION_FILL)

    ws.cell(14, label_col, "Category")
    _style_range(ws, 14, label_col, SECTION_FONT)
    _style_range(ws, 14, value_col, SECTION_FONT)
    if index > 0:
        ws.cell(14, label_col, f"={get_column_letter(first_label_col)}14")

    first_cat = 15
    pulled = amounts or [0.0] * len(categories)
    for i, name in enumerate(categories):
        row = first_cat + i
        if index == 0:
            ws.cell(row, label_col, name)
        else:
            ws.cell(row, label_col, f"={get_column_letter(first_label_col)}{row}")
        value = pulled[i] if i < len(pulled) else 0
        ws.cell(row, value_col, value)
        _style_range(ws, row, label_col, BODY_FONT, align=Alignment(vertical="top"))
        _style_range(ws, row, value_col, AMOUNT_FONT, fmt=CURRENCY_FORMAT)

    total_row = first_cat + len(categories)
    vletter = get_column_letter(value_col)
    ws.cell(total_row, label_col, "TOTAL")
    ws.cell(total_row, value_col, f"=SUM({vletter}{first_cat}:{vletter}{total_row - 1})")
    _style_range(ws, total_row, label_col, TOTAL_COL_FONT, align=Alignment(vertical="center"))
    _style_range(ws, total_row, value_col, TOTAL_COL_FONT, fmt=CURRENCY_FORMAT)

    ws.merge_cells(start_row=4, start_column=spacer_col, end_row=total_row, end_column=spacer_col)
    spacer = ws.cell(4, spacer_col, service_name)
    spacer.font = VERT_FONT
    spacer.fill = SPACER_FILLS[index % len(SPACER_FILLS)]
    spacer.alignment = Alignment(textRotation=90, horizontal="center", vertical="center", wrap_text=True)
    spacer.border = BOX
    return last_row


def build_weekly_report(
    plan: WeekPlan,
    categories: list[str],
    dest: Path,
    amounts_by_service: list[list[float]] | None = None,
) -> Path:
    plan.validate()
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = plan.report_sheet_name()

    ws.row_dimensions[1].height = 15.75
    ws.row_dimensions[2].height = 42.75
    ws.row_dimensions[3].height = 6.75

    last_col = max(3, len(plan.services) * 3)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    title = ws.cell(2, 1, plan.report_title())
    title.font = TITLE_FONT
    title.alignment = Alignment(vertical="top", wrap_text=True)

    first_label_col = 1
    last_row = 26
    for index, service in enumerate(plan.services):
        pulled = amounts_by_service[index] if amounts_by_service and index < len(amounts_by_service) else None
        last_row = write_service_block(ws, service, index, categories, first_label_col, pulled)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A2:{get_column_letter(last_col)}{last_row}"
    ws.sheet_view.showGridLines = False

    wb.save(dest)
    wb.close()
    return dest
