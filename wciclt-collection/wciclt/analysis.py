from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .categories import infer_from_analysis_sheet, write_categories_sheet
from .models import ANALYSIS_ORG, DEFAULT_PREPARED_BY, Service
from .paths import ANALYSIS_TEMPLATE


def _copy_cell(src, dst) -> None:
    dst.value = src.value
    if src.has_style:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)
        dst.alignment = copy(src.alignment)


def copy_sheet(src: Worksheet, dst: Worksheet) -> None:
    for col in range(1, (src.max_column or 1) + 1):
        letter = get_column_letter(col)
        if letter in src.column_dimensions:
            dst.column_dimensions[letter].width = src.column_dimensions[letter].width
            dst.column_dimensions[letter].hidden = src.column_dimensions[letter].hidden
    max_row = src.max_row or 1
    max_col = src.max_column or 1
    for row in range(1, max_row + 1):
        if src.row_dimensions[row].height:
            dst.row_dimensions[row].height = src.row_dimensions[row].height
        for col in range(1, max_col + 1):
            _copy_cell(src.cell(row, col), dst.cell(row, col))
    for merged in src.merged_cells.ranges:
        dst.merge_cells(str(merged))
    dst.page_setup.orientation = src.page_setup.orientation or "landscape"
    dst.page_setup.fitToWidth = 1
    dst.page_setup.fitToHeight = 1
    dst.sheet_properties.pageSetUpPr.fitToPage = True
    if src.sheet_properties.tabColor is not None:
        dst.sheet_properties.tabColor = src.sheet_properties.tabColor
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines


def load_packaged_template_sheet() -> Workbook:
    if not ANALYSIS_TEMPLATE.exists():
        raise FileNotFoundError(f"Missing analysis template: {ANALYSIS_TEMPLATE}")
    return load_workbook(ANALYSIS_TEMPLATE)


def ensure_template_sheet(wb: Workbook) -> Worksheet:
    if "_Template" in wb.sheetnames:
        return wb["_Template"]
    packaged = load_packaged_template_sheet()
    try:
        src = packaged["_Template"] if "_Template" in packaged.sheetnames else packaged[packaged.sheetnames[0]]
        dst = wb.create_sheet("_Template", 0)
        copy_sheet(src, dst)
        dst.sheet_state = "hidden"
        return dst
    finally:
        packaged.close()


def _category_bounds(ws: Worksheet) -> tuple[int, int]:
    start = None
    end = None
    for row in range(30, (ws.max_row or 30) + 1):
        value = ws.cell(row, 2).value
        if value is None:
            continue
        text = str(value).strip().casefold()
        if text == "collection analysis":
            start = row + 2  # blank row then first category (row 35 when heading is 33)
        if text == "total collections analysis":
            end = row
            break
    if start is None:
        start = 35
    if end is None:
        end = start
    return start, end


def apply_categories_to_sheet(ws: Worksheet, categories: list[str]) -> None:
    start, total_row = _category_bounds(ws)
    current = infer_from_analysis_sheet(ws)
    amounts: dict[str, object] = {}
    for offset, name in enumerate(current):
        row = start + offset
        if row >= total_row:
            break
        amounts[name.casefold()] = ws.cell(row, 7).value

    needed = len(categories)
    existing = max(0, total_row - start)
    if needed > existing:
        ws.insert_rows(total_row, needed - existing)
        total_row += needed - existing
    elif needed < existing:
        ws.delete_rows(start + needed, existing - needed)
        total_row -= existing - needed

    sample = ws.cell(start, 2)
    for i, name in enumerate(categories):
        row = start + i
        ws.cell(row, 2).value = name
        if sample.has_style:
            ws.cell(row, 2).font = copy(sample.font)
        prev = amounts.get(name.casefold())
        gcell = ws.cell(row, 7)
        if isinstance(prev, str) and str(prev).startswith("="):
            gcell.value = f"=SUM(D{row}:F{row})"
        elif prev not in (None, "") and name.casefold() in amounts:
            gcell.value = prev
        else:
            gcell.value = 0

    last_cat = start + needed - 1 if needed else start
    ws.cell(total_row, 2).value = "TOTAL COLLECTIONS ANALYSIS"
    for col, letter in ((4, "D"), (5, "E"), (6, "F"), (7, "G")):
        ws.cell(total_row, col).value = f"=SUM({letter}{start - 1}:{letter}{last_cat})"
    # Keep deposit-side labels that sit on the original template rows when they still exist.
    if ws.cell(total_row, 9).value is None or str(ws.cell(total_row, 9).value).strip() == "":
        ws.cell(total_row, 9).value = "TOTAL DEPOSITS"
        ws.cell(total_row, 11).value = "=SUM(K27+K39+K42)"
    balance_row = total_row + 2
    if ws.cell(balance_row, 4).value is None:
        ws.cell(balance_row, 4).value = "Balancing Control"
        ws.cell(balance_row, 9).value = "Balancing Control"
        ws.cell(balance_row, 11).value = f"=SUM(K{total_row}-G{total_row})"


def _unique_tab_name(wb: Workbook, base: str) -> str:
    name = base[:31]
    if name not in wb.sheetnames:
        return name
    for i in range(2, 20):
        candidate = f"{base[:27]} -{i}"
        if candidate not in wb.sheetnames:
            return candidate
    raise ValueError(f"Could not allocate a unique tab name for {base}")


def add_service_tab(wb: Workbook, service: Service, categories: list[str]) -> str:
    template = ensure_template_sheet(wb)
    tab = _unique_tab_name(wb, service.analysis_tab_name())
    ws = wb.copy_worksheet(template)
    ws.title = tab
    ws.sheet_state = "visible"
    if ws["B1"].value in (None, ""):
        ws["B1"] = ANALYSIS_ORG
    ws["B2"] = "COLLECTIONS/DEPOSIT RECON"
    ws["C2"] = service.analysis_c2()
    apply_categories_to_sheet(ws, categories)
    # Ensure prepared-by default if the cell is empty.
    for row in range(50, 62):
        label = ws.cell(row, 2).value
        if label and "PREPARED" in str(label).upper() and not ws.cell(row, 3).value:
            ws.cell(row, 3).value = DEFAULT_PREPARED_BY
    return tab


def create_or_open_analysis(path: Path) -> Workbook:
    if path.exists():
        return load_workbook(path)
    wb = load_packaged_template_sheet()
    if "_Template" in wb.sheetnames:
        wb["_Template"].sheet_state = "hidden"
    write_categories_sheet(wb, [])
    return wb


def add_week_tabs(
    analysis_path: Path,
    services: list[Service],
    categories: list[str],
) -> list[str]:
    analysis_path.parent.mkdir(parents=True, exist_ok=True)
    wb = create_or_open_analysis(analysis_path)
    try:
        write_categories_sheet(wb, categories)
        created: list[str] = []
        for service in services:
            created.append(add_service_tab(wb, service, categories))
        wb.save(analysis_path)
        return created
    finally:
        wb.close()
