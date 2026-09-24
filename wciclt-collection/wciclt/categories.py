from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .models import ANALYSIS_NON_CATEGORY, CATEGORIES_SHEET, DEFAULT_CATEGORIES
from .paths import DEFAULT_CATEGORIES_JSON, ensure_data_dir


THANKSGIVING_NAME = "End of Month Thanks-Giving"


def _norm(name: str) -> str:
    return re.sub(r"\s+", " ", (name or "").strip())


def canonicalize_category(name: str) -> str:
    text = _norm(name)
    key = text.casefold()
    if "thanks" in key and "giving" in key.replace("-", " "):
        return THANKSGIVING_NAME
    if key in {"thanks-giving", "thanksgiving", "thanksgiving april"}:
        return THANKSGIVING_NAME
    return text


def normalize_list(names: list[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for raw in names:
        name = canonicalize_category(raw)
        if not name:
            continue
        key = name.casefold()
        if key in ANALYSIS_NON_CATEGORY or key in seen:
            continue
        seen.add(key)
        out.append(name)
    return out


def load_json_categories(path: Path | None = None) -> list[str]:
    path = path or DEFAULT_CATEGORIES_JSON
    if path.exists():
        data = json.loads(path.read_text(encoding="utf-8"))
        names = data.get("categories", data) if isinstance(data, dict) else data
        return normalize_list(list(names))
    return list(DEFAULT_CATEGORIES)


def save_json_categories(names: list[str], path: Path | None = None) -> Path:
    path = path or DEFAULT_CATEGORIES_JSON
    ensure_data_dir()
    path.write_text(
        json.dumps({"categories": normalize_list(names)}, indent=2) + "\n",
        encoding="utf-8",
    )
    return path


def write_categories_sheet(wb: Workbook, names: list[str]) -> None:
    names = normalize_list(names)
    if CATEGORIES_SHEET in wb.sheetnames:
        ws = wb[CATEGORIES_SHEET]
        wb.remove(ws)
    ws = wb.create_sheet(CATEGORIES_SHEET, 0)
    ws.sheet_state = "hidden"
    ws["A1"] = "Category"
    ws["A1"].font = Font(name="Arial", bold=True, size=12)
    for i, name in enumerate(names, start=2):
        ws.cell(i, 1, name)
    ws.column_dimensions["A"].width = 36


def read_categories_sheet(wb: Workbook) -> list[str]:
    if CATEGORIES_SHEET not in wb.sheetnames:
        return []
    ws = wb[CATEGORIES_SHEET]
    names: list[str] = []
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0]:
            names.append(str(row[0]))
    return normalize_list(names)


def infer_from_analysis_sheet(ws: Worksheet) -> list[str]:
    names: list[str] = []
    started = False
    for row in range(33, (ws.max_row or 33) + 1):
        value = ws.cell(row, 2).value
        if value is None:
            continue
        text = _norm(str(value))
        key = text.casefold()
        if key == "collection analysis":
            started = True
            continue
        if not started:
            continue
        if key == "total collections analysis" or key in ANALYSIS_NON_CATEGORY:
            break
        names.append(text)
    return normalize_list(names)


def infer_from_analysis_workbook(wb: Workbook) -> list[str]:
    from_sheet = read_categories_sheet(wb)
    if from_sheet:
        return from_sheet
    for name in reversed(wb.sheetnames):
        if name.startswith("_"):
            continue
        found = infer_from_analysis_sheet(wb[name])
        if found:
            return found
    return []


def infer_from_report_sheet(ws: Worksheet) -> list[str]:
    names: list[str] = []
    started = False
    for row in range(13, (ws.max_row or 13) + 1):
        value = ws.cell(row, 1).value
        if value is None:
            continue
        text = _norm(str(value))
        if isinstance(value, str) and value.startswith("="):
            continue
        key = text.casefold()
        if key == "category":
            started = True
            continue
        if not started:
            continue
        if key == "total":
            break
        names.append(text)
    return normalize_list(names)


def infer_from_report_workbook(path: Path) -> list[str]:
    wb = load_workbook(path, data_only=False)
    try:
        for name in wb.sheetnames:
            if name.startswith("_"):
                continue
            found = infer_from_report_sheet(wb[name])
            if found:
                return found
    finally:
        wb.close()
    return []


def resolve_categories(
    analysis_wb: Workbook | None = None,
    report_path: Path | None = None,
    explicit: list[str] | None = None,
    json_path: Path | None = None,
) -> list[str]:
    if explicit:
        return normalize_list(explicit)
    if analysis_wb is not None:
        found = infer_from_analysis_workbook(analysis_wb)
        if found:
            return found
    if report_path and Path(report_path).exists():
        found = infer_from_report_workbook(Path(report_path))
        if found:
            return found
    return load_json_categories(json_path)


def merge_unique(*lists: list[str]) -> list[str]:
    return normalize_list([name for group in lists for name in group])
